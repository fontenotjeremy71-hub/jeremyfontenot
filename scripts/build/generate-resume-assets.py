from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "assets" / "resume"
DOCX = OUT / "jeremy-fontenot-resume.docx"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.45)
sec.bottom_margin = Inches(0.45)
sec.left_margin = Inches(0.55)
sec.right_margin = Inches(0.55)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9.2)
styles["Normal"].paragraph_format.space_after = Pt(2.2)

for name, size, bold, color, before, after in [
    ("Resume Name", 20, True, "172033", 0, 1),
    ("Resume Heading", 10.5, True, "1F4E79", 6, 2),
    ("Resume Role", 10, True, "172033", 3, 0),
]:
    st = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.font.name = "Aptos"
    st.font.size = Pt(size)
    st.font.bold = bold
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)

p = doc.add_paragraph(style="Resume Name")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("JEREMY C. FONTENOT")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("IT Service Desk | Windows Systems Administration | Microsoft 365 & Entra").bold = True
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Abbeville, Louisiana  |  jeremyfontenot.online  |  linkedin.com/in/jeremy-fontenot  |  github.com/fontenotjeremy71-hub").font.size = Pt(8.8)

p = doc.add_paragraph()
pPr = p._p.get_or_add_pPr()
pbdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single")
bottom.set(qn("w:sz"), "8")
bottom.set(qn("w:space"), "1")
bottom.set(qn("w:color"), "1F4E79")
pbdr.append(bottom)
pPr.append(pbdr)

def heading(text):
    p = doc.add_paragraph(style="Resume Heading")
    p.add_run(text.upper())

def bullet(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.first_line_indent = Inches(-0.11)
    p.paragraph_format.space_after = Pt(1.3)
    p.add_run("• ").bold = True
    p.add_run(text)

def role(title, company, dates, subtitle):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.columns[0].width = Inches(5.7)
    t.columns[1].width = Inches(1.5)
    p1 = t.cell(0,0).paragraphs[0]
    p1.style = styles["Resume Role"]
    p1.add_run(f"{title} — {company}")
    p2 = t.cell(0,1).paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p2.add_run(dates)
    r.bold = True
    r.font.size = Pt(8.8)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top","left","bottom","right","insideH","insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    t._tbl.tblPr.append(borders)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(subtitle)
    r.italic = True
    r.font.size = Pt(8.9)

heading("Professional Summary")
doc.add_paragraph(
    "IT support professional with enterprise Service Desk experience and hands-on Windows infrastructure administration. "
    "Professional experience includes serving as the sole internal IT resource for a startup, where I built its initial "
    "on-premises Windows environment, deployed Active Directory Domain Services, configured Cloudflare DNS, and established "
    "Microsoft 365 services. Current focus is advancing into Windows Systems Administration through production support, "
    "PowerShell, Microsoft 365/Entra, Intune, virtualization, and an evidence-backed home lab."
)

heading("Core Skills")
for label, text in [
    ("Microsoft", "Active Directory Domain Services (AD DS), Group Policy, Windows Server, Microsoft 365, Microsoft Entra ID, Intune, Windows Autopilot, Azure Arc"),
    ("Administration", "User/group administration, authentication, permissions, workstation integration, DNS, DHCP, account lifecycle, endpoint support"),
    ("Support & ITSM", "Service Desk, incident triage, troubleshooting, escalation, knowledge documentation, Ivanti, ServiceNow"),
    ("Infrastructure", "Hyper-V, SCVMM, Windows Admin Center, Proxmox, pfSense, OpenVPN, Linux"),
    ("Automation", "PowerShell, Microsoft Graph, REST API reporting and administration"),
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.add_run(label + ": ").bold = True
    p.add_run(text)

heading("Professional Experience")
role("Service Desk Technician", "Logicalis", "Apr 2024 – Present", "MSP support for enterprise healthcare and industrial clients")
bullet("Provide enterprise Service Desk support for Summa Health and additional MSP clients, including Windows, application, account, and connectivity troubleshooting.")
bullet("Perform Active Directory account administration, password resets, new-user setup, permissions support, and identity-related troubleshooting within Service Desk scope.")
bullet("Support healthcare applications and workflows including Epic, MyChart, Vocera, and Ivanti ticketing; document resolutions and escalate complex issues to Tier 2/3 teams.")
bullet("Self-reported performance includes an average resolution time of 0.83 days, 85% of incidents resolved without reassignment, and 72% same-day resolution.")

role("Service Desk Representative / Team Lead", "Insight Global", "Oct 2022 – Jan 2024", "Healthesystems service desk support")
bullet("Provided end-user technical troubleshooting, ticket ownership, escalation support, and customer communication in a production support environment.")
bullet("Supported team coordination and technician guidance while maintaining incident ownership and service-quality expectations.")

role("IT Technician", "Completeful Technologies", "Feb 2022 – Jun 2022", "Sole IT administration and startup infrastructure deployment")
bullet("Joined during the launch of a new drop-shipping company and served as the primary internal IT resource responsible for establishing its initial IT environment.")
bullet("Built and configured the on-premises Windows infrastructure, including Windows Server and Active Directory Domain Services for centralized identity, authentication, users, groups, policy, and workstation integration.")
bullet("Configured Cloudflare for authoritative DNS and domain services, and established the company’s Microsoft 365 environment for cloud productivity and identity-related business services.")
bullet("Handled day-to-day IT administration and user support after rollout, covering endpoints, accounts, applications, access, and general technology needs across the organization.")

role("Technical Support", "Hughes Network Systems", "Sep 2021 – Feb 2022", "Customer-facing connectivity and technical support")
bullet("Troubleshot network connectivity, equipment, and service issues while guiding customers through technical resolution steps.")

role("Help Desk", "Remington College", "Dec 2019 – Apr 2021", "Academic IT support")
bullet("Provided user and endpoint support while completing an A.S. in Computer Science - Database Administration.")

heading("Selected Systems Administration Projects")
bullet("Windows LAPS & Advanced Group Policy: schema extension, encrypted password backup, OU delegation, RSOP validation, negative access testing, and forced password rotation.")
bullet("Microsoft Entra Hybrid Identity & Cloud Sync: scoped AD-to-Entra provisioning, Password Hash Sync, source authority, agent health, and provisioning troubleshooting.")
bullet("Windows Autopilot & Intune MDM: user-driven Entra join, Enrollment Status Page, Microsoft 365 Apps, BitLocker, Defender, Secure Boot remediation, and compliance validation.")
bullet("Microsoft 365 Graph Administration Lab: reusable PowerShell/Graph reporting for tenant inventory, users, groups, licensing, directory roles, stale accounts, and Intune devices.")

heading("Education")
p = doc.add_paragraph()
p.add_run("A.S. Computer Science - Database Administration, Remington College").bold = True
p.add_run("  |  Aug 2019 - May 2021  |  GPA 3.84, Honors / Dean's List")
p = doc.add_paragraph()
p.add_run("Information Security & Administration, MyComputerCareer").bold = True
p.add_run("  |  Jun 2021 - Jan 2022  |  GPA 4.00, WCITP")

heading("Certifications")
doc.add_paragraph(
    "Microsoft Azure Fundamentals; CompTIA ITF+, A+, Server+; MTA Windows Server Administration Fundamentals; "
    "MTA Networking Fundamentals; MTA Security Fundamentals; LPI Linux Essentials; Google IT Support; "
    "Cisco Introduction to Cybersecurity; freeCodeCamp Responsive Web Design"
)

OUT.mkdir(parents=True, exist_ok=True)
doc.save(DOCX)
print(DOCX)

# Rebuild trigger for updated downloadable resume assets.
